"""Create a new Excel score chart from a name,score CSV."""
import argparse
import csv
import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def convert(source, destination):
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError('Input and output must be different files')
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Output must use .xlsx')
    records = []
    with source.open(newline='', encoding='utf-8-sig') as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not {'name', 'score'} <= set(reader.fieldnames):
            raise ValueError('CSV requires name and score columns')
        for row in reader:
            name = (row.get('name') or '').strip()
            score = float(row.get('score') or '')
            if not name or len(name) > 200 or not math.isfinite(score):
                raise ValueError('Each row needs a name and a finite numeric score')
            records.append((name, score))
            if len(records) > 10000:
                raise ValueError('Maximum 10000 records')
    if not records:
        raise ValueError('CSV has no records')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Scores'
    sheet.append(['Name', 'Score'])
    for name, score in records:
        sheet.append([name, score])
        sheet.cell(sheet.max_row, 1).data_type = 's'
    sheet.column_dimensions['A'].width = 30
    sheet.column_dimensions['B'].width = 14
    sheet.freeze_panes = 'A2'
    chart = BarChart()
    chart.title = 'Score comparison'
    chart.y_axis.title = 'Score'
    chart.x_axis.title = 'Name'
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=sheet.max_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
    sheet.add_chart(chart, 'D2')
    with destination.open('xb') as stream:
        workbook.save(stream)
    return len(records)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('destination', type=Path)
    args = parser.parse_args()
    try:
        print(f'Exported {convert(args.source, args.destination)} records')
    except (ValueError, OSError) as error:
        parser.exit(2, f'Error: {error}\n')
